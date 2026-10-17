"""Kitchen equipment and exhaust BOQ in the units Indian commercial-kitchen vendors usually quote: sizes in feet and inches (mm alongside),
sheet thickness in SWG, hood and benches per running foot, duct per square foot, airflow in CFM, cooling in TR, volumes in litres.
Taken from the main BOQ (tools/boq.py): sections A and B plus the exhaust lines; gas, plumbing, electrical, civil and fire stay in the main BOQ.
These are common trade conventions (Indian tenders and fabricator listings), not something I could confirm with a Goa vendor: the mm column is there so either works.
Usage: python3 tools/boq_goa.py  ->  boq-goa.html, plans/kitchen-boq-goa.xlsx, plans/kitchen-boq-goa.csv   (PDF: print boq-goa.html from Chrome)
"""
import re, html, csv, boq

MM_IN = 25.4
def ftin(mm):
    """mm -> feet-inches to the nearest half inch, e.g. 800 -> 2'-7½"."""
    half = round(mm / MM_IN * 2); ft, rem = divmod(half, 24); inch, h = divmod(rem, 2)
    return f"{ft}'-{inch}{'½' if h else ''}\"" if ft else f"{inch}{'½' if h else ''}\""
def convert(size):
    """Replace every number in a size string with feet-inches, keeping the words (Ø, long, section...)."""
    def f(m):
        n = m.group(1)
        if m.group(2): return m.group(0)      # a number followed by L or kg is not a length
        return ftin(float(n.replace(',', '')))
    return re.sub(r'(\d[\d,]*)(\s*(?:L|kg)\b)?', f, size)

S = {i[0]: i for sec in boq.S for i in sec[3]}   # item number -> row
def row(no, **kw):
    n, item, size, spec, unit, qty, ref = S[no]
    return dict(no=no, item=kw.get('item', item), mm=kw.get('mm', size), ft=kw.get('ft', convert(size) if size else ''), spec=kw.get('spec', spec),
                unit=kw.get('unit', unit), qty=kw.get('qty', qty), basis=kw.get('basis', ''), bunit=kw.get('bunit', ''), ref=ref)

# Running-foot basis for fabricated pieces (width of the piece in feet), so a per-foot quote can be priced
def rft(mm): return round(mm / 304.8, 1)
SS = 'SS 304'
GOA = [
('1', 'Stainless steel fabrication', f'{SS}. Worktops 18 SWG (1.2 mm), bodies and shelves 19 SWG (1.0 mm), legs 1½" (38 mm) tube, welded and polished joints, adjustable feet. Vendor confirms or corrects the gauge and states it on the quote.', [
 row('A.01', basis=rft(800), bunit='run ft'), row('A.02', basis=rft(800), bunit='run ft'), row('A.03', basis=rft(650), bunit='run ft'), row('A.04', basis=rft(800), bunit='run ft'),
 row('A.05', basis=rft(1400) * 2, bunit='run ft'), row('A.06', basis=rft(450), bunit='run ft'), row('A.07', basis=rft(900), bunit='run ft'), row('A.08', basis=rft(440), bunit='run ft'),
 row('A.09', basis=rft(470), bunit='run ft'), row('A.10', basis=rft(760), bunit='run ft'), row('A.11', basis=rft(650), bunit='run ft'), row('A.12', basis=rft(650), bunit='run ft'),
 row('A.13', basis=rft(400) * 2, bunit='run ft'), row('A.14'), row('A.15', unit='run ft', qty=10.3, spec='SS grating, falls to the grease trap (3.15 m)'),
 row('A.19'), row('A.20'),
]),
('2', 'Cooking, cold and kitchen equipment (branded)', 'State make and model, burner rating in kW (or kcal/hr), litres for each cold unit, and who services it in Goa. Sizes are the space available: do not quote a wider unit.', [
 row('B.01'), row('B.02'), row('B.03'), row('B.04'), row('B.05'), row('B.06'), row('B.07'), row('B.08'), row('B.09'), row('B.10'), row('B.11'),
 row('B.12'), row('B.13'), row('B.14'), row('B.15'), row('B.21'), row('B.22'), row('B.23'),
]),
('3', 'Hood, exhaust and ventilation', 'Exhaust about 2,400 to 3,000 CFM (4,000 to 5,000 m³/h) for the whole hood; make-up air about 80% of it. No roof stack: apartment above and parking behind, so the duct goes out through the back wall, along it and up the right side, with an ESP before the fan. Vendor states the static pressure the fan is selected for.', [
 row('A.16', item='Canopy hood, wall type, with baffle filters', spec=f'{SS}, 20 SWG (0.9 mm); SS baffle filters, grease cup, LED lights; bottom edge 6\'-9" (2.05 m) above the floor; overhang 8" at the open end, 6" over the wok range. Priced per running foot or per hood: state which', basis=11.0, bunit='run ft'),
 row('A.17', item='Exhaust ducting, rectangular', mm='400 × 400 section', ft='16" × 16" section', spec=f'{SS}, 20 SWG (0.9 mm), at 7\'-8" to 9\'-0" above the floor: hood collar, wall opening, along the back wall, round the corner, to the ESP, fan and discharge; brackets, bends, access doors. Sheet area 86 sq ft (about 95 sq ft with bends)', unit='run ft', qty=16.4, basis=95, bunit='sq ft'),
 row('A.18', mm='400 × 400 section', ft='16" × 16" section', spec='About 8\'-6" above the floor, pointing away from the building; bird mesh and rain cowl'),
 row('B.16', item='ESP with carbon odour filter', spec='Outside, 6\'-9" to 9\'-8" above the floor; rated for 2,400 to 3,000 CFM; state kW and filter life'),
 row('B.17', item='Inline exhaust fan', mm='Ø 440', ft='Ø 17½"', spec='2,400 to 3,000 CFM at the static pressure you state (duct + ESP + discharge); belt or direct drive, motor HP stated; anti-vibration mounts, speed control'),
 row('B.18', item='Fresh-air intake fan', spec='High on the back wall (about 10\'-10"); about 1,900 to 2,400 CFM; interlocked with the exhaust'),
 row('B.19', item='Make-up air diffuser, perforated, low velocity', spec='In the kitchen ceiling, about 950 to 1,200 CFM each'),
 row('B.20', item='Make-up air cooling coil with condenser', spec='1.5 TR; condenser on the right side wall at about 9\'-10" above the floor'),
 row('E.03', item='Exhaust fan and ESP control panel', mm='', ft='', spec='With hood-light switch and interlocks (fan, make-up air, gas solenoid, suppression)'),
]),
]

def fmt(q): return f'{q:g}' if isinstance(q, float) else str(q)
def rows():
    for sl, title, note, items in GOA:
        yield ('sec', sl, title, note)
        for k, r in enumerate(items, 1): yield ('item', f'{sl}.{k:02d}', r)

HDR = ['Sl', 'Item / type', 'Size W × D × H (ft-in)', 'Size (mm)', 'Gauge, capacity, specification', 'Unit', 'Qty', 'Per-foot basis qty', 'Basis unit', 'Rate per unit (₹)', 'Rate per basis unit (₹)', 'Amount (₹)', 'Make / model / remarks', 'Main BOQ no.']

def write_csv(path):
    with open(path, 'w', newline='') as f:
        w = csv.writer(f); w.writerow(HDR)
        for r in rows():
            if r[0] == 'sec': w.writerow([r[1], r[2], '', '', r[3]] + [''] * 9)
            else:
                d = r[2]; w.writerow([r[1], d['item'], d['ft'], d['mm'], d['spec'], d['unit'], fmt(d['qty']), fmt(d['basis']) if d['basis'] != '' else '', d['bunit'], '', '', '', '', d['no']])

def write_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    wb = Workbook(); ws = wb.active; ws.title = 'BOQ'
    thin = Side(style='thin', color='BBBBBB'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill('solid', fgColor='2A2420'); band = PatternFill('solid', fgColor='EFE6D3'); inp = PatternFill('solid', fgColor='FFF9E0')
    NC = len(HDR)
    ws['A1'] = 'BOQ · kitchen equipment and exhaust · Betalbatim pub'; ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = 'Shop 4, Tulip Residency, Gandaulim, Salcete, South Goa. Sizes in feet and inches with mm alongside; sheet in SWG; airflow in CFM; cooling in TR. Gas, plumbing, electrical, civil and fire work are in the main BOQ.'
    ws['A3'] = 'Fill the yellow cells. Enter a rate per unit, OR a rate per running foot or per sq ft in the basis column: the amount takes whichever you fill (per unit first). Rates ex-GST; freight, installation and GST at the bottom.'
    for r in (2, 3): ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC); ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[r].height = 30
    for i, h in enumerate(HDR, 1):
        c = ws.cell(5, i, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = head; c.alignment = Alignment(vertical='center', wrap_text=True); c.border = box
    ws.row_dimensions[5].height = 42
    r = 6; subs = []
    for sl, title, note, items in GOA:
        ws.cell(r, 1, sl); ws.cell(r, 2, title); ws.cell(r, 5, note)
        for i in range(1, NC + 1): c = ws.cell(r, i); c.fill = band; c.font = Font(bold=True); c.border = box; c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 62; r += 1; first = r
        for k, d in enumerate(items, 1):
            vals = [f'{sl}.{k:02d}', d['item'], d['ft'], d['mm'], d['spec'], d['unit'], d['qty'], d['basis'] if d['basis'] != '' else None, d['bunit'] or None, None, None,
                    f'=IF(J{r}<>"",F{r}*J{r},IF(AND(K{r}<>"",H{r}<>""),H{r}*K{r},""))', None, d['no']]
            vals[5], vals[6] = d['unit'], d['qty']
            vals = [vals[0], vals[1], vals[2], vals[3], vals[4], vals[5], vals[6], vals[7], vals[8], vals[9], vals[10], vals[11], vals[12], vals[13]]
            for i, v in enumerate(vals, 1):
                c = ws.cell(r, i, v); c.border = box; c.alignment = Alignment(wrap_text=True, vertical='top')
            # formula uses qty in column G, rate per unit J: fix to the real columns
            ws.cell(r, 12, f'=IF(J{r}<>"",G{r}*J{r},IF(AND(K{r}<>"",H{r}<>""),H{r}*K{r},""))')
            for i in (10, 11, 13): ws.cell(r, i).fill = inp
            for i in (10, 11, 12): ws.cell(r, i).number_format = '#,##0.00'
            ws.cell(r, 7).number_format = '0.0#'; ws.cell(r, 8).number_format = '0.0#'
            r += 1
        ws.cell(r, 2, f'Subtotal, section {sl}'); ws.cell(r, 12, f'=SUM(L{first}:L{r-1})'); subs.append(r)
        for i in range(1, NC + 1): c = ws.cell(r, i); c.font = Font(bold=True); c.border = box
        ws.cell(r, 12).number_format = '#,##0.00'; r += 2
    ws.cell(r, 2, 'Total of all sections (ex-GST)'); ws.cell(r, 12, '=' + '+'.join(f'L{x}' for x in subs)); tot = r; r += 1
    ws.cell(r, 2, 'Freight and unloading to Betalbatim'); ws.cell(r, 12).fill = inp; fr = r; r += 1
    ws.cell(r, 2, 'Installation and connections not included above'); ws.cell(r, 12).fill = inp; ins = r; r += 1
    ws.cell(r, 2, 'GST rate (%)'); ws.cell(r, 12, 18).fill = inp; gr = r; r += 1
    ws.cell(r, 2, 'GST'); ws.cell(r, 12, f'=ROUND((L{tot}+L{fr}+L{ins})*L{gr}/100,2)'); gs = r; r += 1
    ws.cell(r, 2, 'Grand total including GST'); ws.cell(r, 12, f'=L{tot}+L{fr}+L{ins}+L{gs}'); gt = r
    for x in range(tot, r + 1):
        ws.cell(x, 2).font = Font(bold=True); ws.cell(x, 12).number_format = '#,##0.00'; ws.cell(x, 12).font = Font(bold=True)
    ws.cell(gt, 2).fill = band; ws.cell(gt, 12).fill = band
    for c, w in zip('ABCDEFGHIJKLMN', (6, 34, 20, 20, 58, 7, 7, 10, 8, 13, 13, 14, 26, 9)): ws.column_dimensions[c].width = w
    ws.freeze_panes = 'C6'; ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = '5:5'
    n = wb.create_sheet('Notes')
    for i, t in enumerate([
        'Units used in this BOQ',
        'Length: feet and inches (nearest ½ inch) with the millimetres alongside. Sheet thickness: SWG with the millimetres. Hood and benches can be priced per running foot, duct per square foot, or per unit.',
        'Airflow: CFM (1 m³/h = 0.589 CFM). Cooling: TR. Cold units: litres. Cylinders: kg. Burners: kW (or kcal/hr, say which).',
        'These are the conventions commonly used in Indian commercial-kitchen quotes; if your vendor works in a different unit, use the millimetre column.',
        'Every quote must state: steel grade and SWG, price ex-GST with the GST rate, delivery, installation, lead time, payment schedule, warranty and who services it in Goa, validity.',
        'Concept drawings for briefing suppliers, not construction drawings. Duct length is an estimate: the vendor measures on site.']):
        n.cell(i + 1, 1, t).alignment = Alignment(wrap_text=True, vertical='top')
    n['A1'].font = Font(bold=True, size=13); n.column_dimensions['A'].width = 130
    wb.save(path)

def page():
    esc = html.escape; body = []; n = 0
    for r in rows():
        if r[0] == 'sec': body.append(f'<tr class="sec"><th>{r[1]}</th><th colspan="9">{esc(r[2])}<small>{esc(r[3])}</small></th></tr>')
        else:
            d = r[2]; n += 1
            b = f'{fmt(d["basis"])} {d["bunit"]}' if d['basis'] != '' else ''
            body.append(f'<tr><td class="no">{r[1]}</td><td><b>{esc(d["item"])}</b></td><td class="sz">{esc(d["ft"])}</td><td class="sz">{esc(d["mm"])}</td><td>{esc(d["spec"])}</td><td>{d["unit"]}</td><td class="n">{fmt(d["qty"])}</td><td class="n">{b}</td><td class="blank"></td><td class="ref">{d["no"]}</td></tr>')
    return f'''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover"><title>Kitchen Equipment and Exhaust BOQ · Betalbatim Pub</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:opsz,wght@12..96,700&family=IBM+Plex+Sans:wght@400;500;600;700&family=IBM+Plex+Mono:wght@500&display=swap">
<style>
/* Layout: drafting-paper page; the table scrolls sideways on a phone and prints on A4 landscape with a repeating header */
:root{{--paper:#fbf7ef;--ink:#2a2420;--muted:#6a5d4e;--rule:#d9cbb3;--accent:#b23a1e;--chip:#fffdf8;--display:'Bricolage Grotesque','Segoe UI',sans-serif;--body:'IBM Plex Sans','Helvetica Neue',Arial,sans-serif;--mono:'IBM Plex Mono',ui-monospace,Menlo,monospace;color-scheme:light}}
@page{{size:A4 landscape;margin:10mm}}
body{{background:var(--paper);color:var(--ink);font-family:var(--body);margin:0;line-height:1.45}}
.wrap{{max-width:1240px;margin:0 auto;padding:22px 16px 50px;display:flex;flex-direction:column;gap:16px}}
.eyebrow{{font:500 11px/1.4 var(--mono);letter-spacing:.08em;text-transform:uppercase;color:var(--accent);margin:0}}
h1{{font:700 clamp(24px,4vw,34px)/1.05 var(--display);margin:4px 0 6px}}
.lede{{margin:0;max-width:84ch;color:var(--muted)}}
.bar{{display:flex;flex-wrap:wrap;gap:8px}} .bar a{{font:600 13px var(--body);color:var(--ink);background:var(--chip);border:1px solid var(--rule);border-radius:8px;padding:9px 12px;text-decoration:none}} .bar a.pri{{background:var(--ink);color:var(--paper);border-color:var(--ink)}}
.tw{{overflow-x:auto;border:1px solid var(--rule);border-radius:8px;background:var(--chip)}}
table{{border-collapse:collapse;width:100%;min-width:1080px;font-size:13px}}
th,td{{padding:6px 9px;border-top:1px solid var(--rule);vertical-align:top;text-align:left}}
thead th{{background:var(--ink);color:var(--paper);font:600 11.5px var(--mono);border-top:0;white-space:nowrap}}
tr.sec th{{background:#efe6d3;font:700 13px var(--body)}} tr.sec small{{display:block;font:400 12px var(--body);color:var(--muted);margin-top:2px}}
td.no{{font:600 12px var(--mono);white-space:nowrap}} td.sz{{font:500 12px var(--mono);white-space:nowrap}} td.n{{text-align:right;font-family:var(--mono);white-space:nowrap}} td.ref{{color:var(--muted);font-size:12px}}
td.blank{{background:#fffdf0;min-width:80px}}
a{{color:var(--accent)}} ul{{margin:0;padding-left:20px;font-size:14px}} li{{margin:3px 0}}
@media print{{.bar{{display:none}} body{{background:#fff}} .wrap{{padding:0}} .tw{{border:0;overflow:visible}} table{{min-width:0;font-size:8pt}} th,td{{padding:3px 5px}} td.blank{{min-width:0;width:9%}} td.sz{{white-space:normal}} td.n{{white-space:normal}} tr.sec small{{font-size:7.5pt}} thead{{display:table-header-group}} tr{{break-inside:avoid}}}}
</style></head><body><div class="wrap">
<header><p class="eyebrow">Simplex · Betalbatim pub · kitchen</p><h1>BOQ: kitchen equipment and exhaust</h1>
<p class="lede">{n} lines in 3 sections: stainless steel fabrication, cooking and cold equipment, and the hood, exhaust and ventilation. Sizes are in feet and inches with millimetres alongside, sheet in SWG, hood and benches by the running foot, duct by the square foot, airflow in CFM, cooling in TR, cold units in litres. These are the units Indian commercial-kitchen quotes usually use; I could not confirm them with a Goa vendor, so the millimetre column is there if they prefer it. Gas, plumbing, electrical, civil and fire work are in the <a href="boq.html">main BOQ</a>. Concept drawings for briefing suppliers, not construction drawings.</p></header>
<div class="bar"><a class="pri" href="plans/kitchen-boq-goa.xlsx">Excel (per unit, per foot or per sq ft)</a><a href="plans/kitchen-boq-goa.pdf" target="_blank" rel="noopener">PDF, A4</a><a href="plans/kitchen-boq-goa.csv">CSV</a><a href="boq.html">Main BOQ (metric)</a><a href="boq-drawing.html">Drawing + BOQ in mm (A3 PDF)</a><a href="kitchen.html">Kitchen sheets</a></div>
<div class="tw"><table><thead><tr><th>Sl</th><th>Item / type</th><th>Size W × D × H (ft-in)</th><th>Size (mm)</th><th>Gauge, capacity, specification</th><th>Unit</th><th>Qty</th><th>Per-foot basis</th><th>Rate ₹</th><th>Main BOQ</th></tr></thead><tbody>{''.join(body)}</tbody></table></div>
<section><b>Every quote must state</b><ul>
<li>Steel grade and SWG for each part; make and model, kW or litres for every branded item; the static pressure the exhaust fan is selected for.</li>
<li>Price ex-GST with the GST rate separate; delivery to Betalbatim and unloading; installation and connections included or extra; whether the hood and benches are priced per unit or per running foot, and the duct per square foot.</li>
<li>Lead time, payment schedule, warranty and who services it in Goa, and how long the quote is valid.</li></ul></section>
</div></body></html>
'''

if __name__ == '__main__':
    write_csv('plans/kitchen-boq-goa.csv'); write_xlsx('plans/kitchen-boq-goa.xlsx')
    open('boq-goa.html', 'w').write(page())
    print(sum(len(i) for *_, i in GOA), 'lines; boq-goa.html, plans/kitchen-boq-goa.xlsx, .csv written')
