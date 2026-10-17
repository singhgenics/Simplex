"""Bill of quantities (BOQ) for the kitchen, in the usual form: numbered items by trade, specification, unit, quantity, rate, amount.
Quantities are taken from the model coordinates (same source as the plan and the A3 sheets) and are for pricing; the contractor measures on site.
Rates are left blank for the suppliers; the Excel file multiplies and totals them.
Usage: python3 tools/boq.py  ->  boq.html, plans/kitchen-boq.xlsx, plans/kitchen-boq.csv   (PDF: print boq.html from Chrome, see tools/README.md)
"""
import html, csv, sys

# (item, description and specification, unit, qty, drawing reference / note)
S = [
('A', 'Stainless steel fabrication (SS 304 throughout)', 'Worktops 1.2 mm or thicker, bodies 1.0 mm, legs 38 mm tube, welded and polished joints, adjustable feet. Supplier confirms or corrects the gauges.', [
 ('A.01', 'Non-veg prep bench', '800 × 650 × 900', 'Open undershelf for the board rack, 100 mm back upstand', 'nos', 1, 'Sheet 1'),
 ('A.02', 'Dumpling and veg prep bench', '800 × 650 × 900', 'Top over the under-counter freezer (B.07); two wall shelves over', 'nos', 1, 'Sheet 1'),
 ('A.03', 'Pass counter', '650 × 650 × 900', 'Heat-lamp bracket; hatch ledge 700 mm through the wall', 'nos', 1, 'Sheet 1'),
 ('A.04', 'Dish sink, 2 bowls', '800 × 650 × 900', 'Pre-rinse spray bracket, undershelf, legs', 'nos', 1, 'Sheet 1'),
 ('A.05', 'Wall shelf over the dish line', '1,400 × 350', 'One plain, one slatted for bamboo steamer baskets', 'nos', 2, 'Sheet 1'),
 ('A.06', 'Veg and herb wash sink', '450 × 500 × 900', 'Single bowl, legs, undershelf', 'nos', 1, 'Sheet 1'),
 ('A.07', 'Sauce and oil wall shelf', '900 × 300', 'Back-up stock, fixed at 1.45 m', 'nos', 1, 'Sheet 1'),
 ('A.08', 'Wok seasoning rail', '440 × 220', 'At 1.0 m; 8 lidded pots and a rack for 8 squeeze bottles', 'set', 1, 'Sheet 1'),
 ('A.09', 'Wok tool rail', '470 long', 'At 1.4 m with 4 hooks', 'nos', 1, 'Sheet 1'),
 ('A.10', 'Chilled mise-en-place rail', '760 × 260 × 220', 'Counter-top unit for 5 pans', 'nos', 1, 'Sheet 1'),
 ('A.11', 'Dip and garnish rail', '650 × 120', '5 inserts, on the pass', 'nos', 1, 'Sheet 1'),
 ('A.12', 'Ticket rail', '650 long', 'Above the pass hatch', 'nos', 1, 'Sheet 1'),
 ('A.13', 'Magnetic knife strip', '400 long', '', 'nos', 2, 'Sheet 1'),
 ('A.14', 'Cutting-board rack', '280 × 420 × 420', 'Slots for 4 colour-coded boards, under the non-veg bench', 'nos', 1, 'Sheet 1'),
 ('A.15', 'Floor channel drain', '3,150 × 80', 'SS grating, falls to the grease trap', 'm', 3.15, 'Sheet 1'),
 ('A.16', 'Canopy hood', '3,350 × 950 × 550', 'Bottom edge at 2.05 m; baffle filters, grease cup, LED lights; overhang 0.20 m at the open end, 0.15 m over the wok range', 'nos', 1, 'Sheet 2'),
 ('A.17', 'Exhaust ductwork', '400 × 400 section', 'SS at 2.35 to 2.75 m: hood collar, wall penetration, along the back wall, round the corner, to the ESP, fan and discharge; brackets, bends and access doors', 'm', 5.0, 'Sheet 2 · measure on site'),
 ('A.18', 'Discharge elbow', '400 × 400 section', 'About 2.6 m high, pointing away from the building; bird mesh and rain cowl', 'nos', 1, 'Sheet 2'),
 ('A.19', 'Storage deck racking', '450 × 2,050 · 1,900 × 450 · 800 × 450', 'SS or galvanised, 3 tiers, one rack of each size', 'nos', 3, 'Model'),
 ('A.20', 'Ceiling hatch frame with fold-down ladder', '600 × 800 opening', 'For light goods; heavy stock goes on the hoist (B.21)', 'nos', 1, 'Sheet 1'),
]),
('B', 'Cooking, cold and kitchen equipment (branded)', 'State make and model, kW or litres, and who services it in Goa. Sizes are the space available; do not quote a wider unit.', [
 ('B.01', 'Chinese wok range, 2 burners', '1,200 × 800 × 800', 'LPG; state burner kW. 750 to 900 mm deep is acceptable (aisle is 0.95 m at the range)', 'nos', 1, 'Sheet 1'),
 ('B.02', 'Dim sum steamer, 2 holes', '600 × 700 × 800', 'LPG; with water top-up', 'nos', 1, 'Sheet 1'),
 ('B.03', 'Twin-tank fryer', '600 × 700 × 900', 'LPG; two tanks (veg and non-veg oil)', 'nos', 1, 'Sheet 1'),
 ('B.04', 'Stock range, 2 burners', '600 × 700 × 850', 'LPG; for 20 L stock pots', 'nos', 1, 'Sheet 1'),
 ('B.05', 'Reach-in, 2 doors, dual temperature', '950 × 750 × 2,000', 'About 800 L: fridge door at the front, freezer door at the back, doors open into the aisle; state litres of each side', 'nos', 1, 'Sheet 1'),
 ('B.06', 'Under-counter fridge (pass)', '600 × 550 × 700', 'Sauce bases and cut veg; state litres', 'nos', 1, 'Sheet 1'),
 ('B.07', 'Under-counter freezer, 2 doors', 'fits under 800 × 650', 'State litres (target about 650 L of freezer in all with the reach-in)', 'nos', 1, 'Sheet 1'),
 ('B.08', 'Under-counter dishwasher', '600 × 650 × 900', 'With drainboard and its own booster heater', 'nos', 1, 'Sheet 1'),
 ('B.09', 'Rice cooker', 'Ø 260', '5 L; one for jasmine, one for sticky rice', 'nos', 2, 'Sheet 1'),
 ('B.10', 'Mixer-grinder', '', 'Commercial, for sauces and batter', 'nos', 1, 'Sheet 1'),
 ('B.11', 'Kitchen blowtorch', '', '', 'nos', 1, 'Sheet 1'),
 ('B.12', 'Storage water heater', 'Ø 330 × 500', '35 L, 3 kW, 70 °C, high on the left wall over the dish line', 'nos', 1, 'Sheet 2'),
 ('B.13', 'Water heater for a hand-wash basin', '', '6 L, under each basin', 'nos', 2, 'Sheet 2'),
 ('B.14', 'Insect killer', '500 × 180', 'Over the kitchen door', 'nos', 1, 'Sheet 1'),
 ('B.15', 'Heat lamp over the pass', '', 'On its own switch', 'nos', 2, 'Sheet 1'),
 ('B.16', 'ESP with carbon odour filter', '500 × 1,000 × 900', 'Outside at 2.05 to 2.95 m', 'nos', 1, 'Sheet 2'),
 ('B.17', 'Inline exhaust fan', 'Ø 440', '4,000 to 5,000 m³/h; anti-vibration mounts, speed control', 'nos', 1, 'Sheet 2'),
 ('B.18', 'Fresh-air intake fan', '500 × 480', 'High on the back wall; interlocked with the exhaust', 'nos', 1, 'Sheet 2'),
 ('B.19', 'Make-up air diffuser', '600 × 500', 'Perforated, low velocity, in the kitchen ceiling', 'nos', 2, 'Sheet 2'),
 ('B.20', 'Make-up air cooling coil with condenser', 'condenser 800 × 550 × 300', '1.5 TR; condenser on the right side wall at about 3.0 m', 'set', 1, 'Sheet 2'),
 ('B.21', 'Goods hoist', '', 'Electric, 200 kg, over the ceiling hatch, with limit switch', 'nos', 1, 'Sheet 2'),
 ('B.22', 'Smallware', '', 'To the chef\'s list: carbon-steel woks, ladle, spatula, spider, wok brush, pans, food pans, knives, boards, stone mortar', 'lot', 1, 'Chef'),
 ('B.23', 'Bamboo steamer stack', 'Ø 260', '3 baskets and a lid, for the drying shelf (the steamer comes with its own baskets, B.02)', 'nos', 3, 'Sheet 1'),
]),
('C', 'Gas installation (licensed contractor, for the fire NOC)', 'Pressure test and a written certificate are part of the work.', [
 ('C.01', 'LPG cage', '350 × 1,700 × 1,600', 'Locked and ventilated, on the right strip', 'nos', 1, 'Sheet 2'),
 ('C.02', 'LPG cylinder with changeover manifold and regulator', '19 kg each', '4 on line and 1 spare, 95 kg in all', 'nos', 5, 'Sheet 2'),
 ('C.03', 'Copper gas pipe', 'size by contractor', 'With clamps and sleeves: along the right outside wall at 1.3 m, in at the back corner, drop to the manifold at 0.4 m, along the back of the hot line', 'm', 9.0, 'Sheet 2 · measure on site'),
 ('C.04', 'Gas solenoid valve', '', 'Weatherproof, outside on the entry; closes with the exhaust fan or the hood suppression', 'nos', 1, 'Sheet 2'),
 ('C.05', 'Ball valve, one per appliance', '', 'On the manifold', 'nos', 6, 'Sheet 2'),
 ('C.06', 'Flexible gas hose, certified', '', 'One per appliance', 'nos', 6, 'Sheet 2'),
 ('C.07', 'LPG leak detector', '', 'With cut-off signal, low on the right wall', 'nos', 1, 'Sheet 2'),
 ('C.08', 'Pressure test, commissioning and certificate', '', '', 'lot', 1, ''),
]),
('D', 'Plumbing and drainage', 'Sizes and fittings by the plumber. All kitchen drains go to the grease trap, then the grey-water line (never the biodigester).', [
 ('D.01', 'Water sump, buried', '1,300 × 1,000 × 800', '1,000 L, with lid and tanker inlet, under the garden', 'nos', 1, 'Plan'),
 ('D.02', 'Pressure pump set with sediment filter and softener', '', '0.5 HP pump; fits in the water cabinet (D.03)', 'set', 1, 'Sheet 2'),
 ('D.03', 'Water cabinet', '450 × 800 × 1,500', 'Ventilated, on the right strip; fabrication included', 'nos', 1, 'Sheet 2'),
 ('D.04', 'Cold and hot water pipe with fittings and isolation valves', 'CPVC or PPR', 'Cabinet to the kitchen along the strip, and the hot-water run to the sink and dishwasher', 'lot', 1, 'Sheet 2 · about 20 m'),
 ('D.05', 'Hand-wash basin, wall-mounted', '330 × 400 and 250 × 250', 'Hands-free tap and soap dispenser: one at the door, one at the prep corner', 'nos', 2, 'Sheet 1'),
 ('D.06', 'Swan-neck wok tap', '', 'Wall-mounted, tall, at the wok range', 'nos', 1, 'Sheet 2'),
 ('D.07', 'Mixer taps for the veg sink and the dish sink', '', 'Tall mixer tap for the veg sink; sink mixer with pre-rinse spray for the dish sink', 'set', 1, 'Sheet 2'),
 ('D.08', 'Pot filler', '', 'Wall-mounted swing arm, cold, over the stock range', 'nos', 1, 'Sheet 2'),
 ('D.09', 'Steamer water feed with isolation valve', '', 'At the back wall', 'nos', 1, 'Sheet 2'),
 ('D.10', 'Grease trap', '450 × 450', 'Outside back corner, with cover', 'nos', 1, 'Sheet 2'),
 ('D.11', 'Drain pipe uPVC with fittings', '100 mm', 'Floor channel, sinks, hand-wash basins and dishwasher to the grease trap and grey-water line', 'm', 12.0, 'Sheet 2 · measure on site'),
 ('D.12', 'Floor trap with grating', '', 'SS', 'nos', 2, 'Sheet 2'),
]),
('E', 'Electrical', 'Three-phase supply about 15 to 20 kW. The fridges, exhaust fan, ESP, lights and water pump are on the power backup. Point numbers are on sheet 2.', [
 ('E.01', 'Socket or switch points 1 to 14', '', 'As marked on sheet 2 (fridge, freezer, rice cookers, mise-en-place rail, heat lamps, dishwasher, insect killer, hood and panel, pass fridge, grinder, heater isolator, hand-wash heaters, hoist), with their own MCBs where marked', 'nos', 14, 'Sheet 2'),
 ('E.02', 'Kitchen sub-distribution board', '', 'With RCCB and MCBs, earthing and labelling', 'nos', 1, 'Sheet 1'),
 ('E.03', 'Exhaust fan and ESP control panel', '', 'With hood-light switch and interlocks (fan, make-up air, gas solenoid, suppression)', 'nos', 1, 'Sheet 2'),
 ('E.04', 'LED hood light', '300 × 200', 'Vapour-proof', 'nos', 3, 'Sheet 2'),
 ('E.05', 'LED ceiling panel', '600 × 600', '', 'nos', 3, 'Model'),
 ('E.06', 'Cable, conduit, trunking and fixings', '', 'For all kitchen points', 'lot', 1, ''),
 ('E.07', 'Emergency light and EXIT sign', '', 'At the rear door', 'set', 1, 'Sheet 1'),
 ('E.08', 'Testing and commissioning of all circuits', '', '', 'lot', 1, ''),
]),
('F', 'Civil and finishes', 'Quantities from the 4.9 × 2.4 m kitchen; add wastage as usual. Ceiling must be non-combustible (no PVC).', [
 ('F.01', 'Wall tiles, white glazed', '75 mm tile', 'To 2.1 m on the left wall and the partition, 0.85 to 2.1 m on the back and right walls; grout and corner trims', 'm²', 22.4, 'Model'),
 ('F.02', 'Anti-slip epoxy floor, grey', '', 'With falls to the channel drain', 'm²', 11.8, 'Model'),
 ('F.03', 'Cement-board ceiling at 2.70 m', '', 'Non-combustible, with frame and white paint', 'm²', 11.8, 'Model'),
 ('F.04', 'Double-acting kitchen door', '800 × 2,050', 'Stainless, spring pivots, vision panel, push plates and kick plate', 'nos', 1, 'Sheet 1'),
 ('F.05', 'Rear exit door', '800 × 2,100', 'Opening outward; closer, rubber sweep, PVC strip curtain and panic bar', 'set', 1, 'Sheet 1'),
 ('F.06', 'Cutting and making good', '400 × 400 duct opening', 'Duct opening in the back wall, gas and water sleeves, fan and intake openings', 'lot', 1, 'Sheet 2'),
 ('F.07', 'Storage deck over the kitchen', '4,900 × 2,400', 'Plywood on steel joists, rated for storage only', 'm²', 11.8, 'Model'),
]),
('G', 'Fire and safety', 'Needed for the fire NOC.', [
 ('G.01', 'Wet-chemical hood suppression', '', 'Nozzles over each appliance, manual pull and automatic gas cut-off', 'set', 1, 'Sheet 2'),
 ('G.02', 'Extinguisher, wet chemical (Class K)', '6 L', 'On the back wall at the end of the hot line', 'nos', 1, 'Sheet 1'),
 ('G.03', 'Extinguisher, CO₂', '4.5 kg', 'Beside the electrical board at the kitchen door', 'nos', 1, 'Model'),
 ('G.04', 'Fire blanket in a wall box', '', '', 'nos', 1, 'Sheet 1'),
 ('G.05', 'First-aid box', '', 'On the wall', 'nos', 1, 'Model'),
 ('G.06', 'Heat detector', '', 'On the kitchen ceiling, wired to the alarm panel', 'nos', 1, 'Sheet 2'),
]),
('H', 'Outside, bins and housekeeping', '', [
 ('H.01', 'Bin store', '450 × 900 × 1,000', 'Timber slats, on the right strip', 'nos', 1, 'Plan'),
 ('H.02', 'Wheeled bin', '60 L', 'Wet, dry and glass', 'nos', 3, 'Plan'),
 ('H.03', 'Foot-pedal bin', 'Ø 200 × 600', 'Wet under the veg sink, dry by the rear door', 'nos', 2, 'Sheet 1'),
 ('H.04', 'Steel bollard', 'Ø 130 × 900', 'On the edge of the strip', 'nos', 6, 'Plan'),
 ('H.05', 'Apron hook rail', '3 hooks', 'Above the prep-corner hand-wash', 'nos', 1, 'Sheet 1'),
 ('H.06', 'Dirty-linen bag with lid; used-oil can with spill tray', 'oil can 20 L', '', 'set', 1, 'Sheet 1'),
]),
]

def rows():
    for sec, title, note, items in S:
        yield ('sec', sec, title, note)
        for it in items: yield ('item',) + it

def counts():
    n = sum(len(i) for _, _, _, i in S); return n

def write_csv(path):
    with open(path, 'w', newline='') as f:
        w = csv.writer(f); w.writerow(['Item', 'Item / type', 'Size W × D × H (mm)', 'Capacity, rating, specification', 'Unit', 'Qty', 'Rate (INR)', 'Amount (INR)', 'Make / model / remarks', 'Drawing'])
        for r in rows():
            if r[0] == 'sec': w.writerow([r[1], r[2], '', r[3], '', '', '', '', '', ''])
            else: w.writerow([r[1], r[2], r[3], r[4], r[5], r[6], '', '', '', r[7]])

def write_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    wb = Workbook(); ws = wb.active; ws.title = 'BOQ'
    thin = Side(style='thin', color='BBBBBB'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill('solid', fgColor='2A2420'); band = PatternFill('solid', fgColor='EFE6D3'); inp = PatternFill('solid', fgColor='FFF9E0')
    NC = 10
    ws['A1'] = 'Bill of quantities · kitchen · Betalbatim pub'; ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = 'Shop 4, Tulip Residency, Gandaulim, Salcete, South Goa. Kitchen 4.9 × 2.4 m, one chef and one helper. Quantities and sizes come from the model and the two A3 sheets (kitchen.html); measure on site before ordering.'
    ws['A3'] = 'Fill the yellow cells (rate, make/model, remarks). Amounts and totals calculate. Rates ex-GST, itemised; freight, installation and the GST rate go in the lines at the bottom. Sizes are width × depth × height in mm, the space available.'
    for r in (2, 3): ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC); ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[r].height = 30
    hdr = ['Item', 'Item / type', 'Size W × D × H (mm)', 'Capacity, rating, specification', 'Unit', 'Qty', 'Rate (₹)', 'Amount (₹)', 'Make / model / remarks', 'Drawing']
    for i, h in enumerate(hdr, 1):
        c = ws.cell(5, i, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = head; c.alignment = Alignment(vertical='center', wrap_text=True); c.border = box
    r = 6; subs = []
    for sec, title, note, items in S:
        ws.cell(r, 1, sec); ws.cell(r, 2, title); ws.cell(r, 4, note)
        for i in range(1, NC + 1): c = ws.cell(r, i); c.fill = band; c.font = Font(bold=True); c.border = box; c.alignment = Alignment(wrap_text=True, vertical='top')
        r += 1; first = r
        for no, item, size, spec, unit, qty, ref in items:
            vals = [no, item, size, spec, unit, qty, None, f'=IF(G{r}="","",F{r}*G{r})', None, ref]
            for i, v in enumerate(vals, 1):
                c = ws.cell(r, i, v); c.border = box; c.alignment = Alignment(wrap_text=True, vertical='top')
            for i in (7, 9): ws.cell(r, i).fill = inp
            ws.cell(r, 7).number_format = '#,##0.00'; ws.cell(r, 8).number_format = '#,##0.00'; ws.cell(r, 6).number_format = '0.0#'
            r += 1
        ws.cell(r, 2, f'Subtotal, section {sec}'); ws.cell(r, 8, f'=SUM(H{first}:H{r-1})'); subs.append(r)
        for i in range(1, NC + 1): c = ws.cell(r, i); c.font = Font(bold=True); c.border = box
        ws.cell(r, 8).number_format = '#,##0.00'; r += 2
    ws.cell(r, 2, 'Total of all sections (ex-GST)'); ws.cell(r, 8, '=' + '+'.join(f'H{x}' for x in subs)); tot = r; r += 1
    ws.cell(r, 2, 'Freight and unloading to Betalbatim'); ws.cell(r, 8).fill = inp; fr = r; r += 1
    ws.cell(r, 2, 'Installation and connections not included above'); ws.cell(r, 8).fill = inp; ins = r; r += 1
    ws.cell(r, 2, 'GST rate (%)'); ws.cell(r, 8, 18).fill = inp; gr = r; r += 1
    ws.cell(r, 2, 'GST'); ws.cell(r, 8, f'=ROUND((H{tot}+H{fr}+H{ins})*H{gr}/100,2)'); gs = r; r += 1
    ws.cell(r, 2, 'Grand total including GST'); ws.cell(r, 8, f'=H{tot}+H{fr}+H{ins}+H{gs}'); gt = r
    for x in range(tot, r + 1):
        ws.cell(x, 2).font = Font(bold=True); ws.cell(x, 8).number_format = '#,##0.00'; ws.cell(x, 8).font = Font(bold=True)
    ws.cell(gt, 2).fill = band; ws.cell(gt, 8).fill = band
    for c, w in zip('ABCDEFGHIJ', (7, 36, 22, 54, 6, 7, 12, 14, 28, 16)): ws.column_dimensions[c].width = w
    ws.freeze_panes = 'C6'; ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = '5:5'
    n = wb.create_sheet('Notes')
    for i, t in enumerate([
        'How to use this BOQ',
        '1. Send the same BOQ and the two A3 sheets to three suppliers per group: A (fabrication) to Gujarat, Mumbai and Goa fabricators; B to Goa dealers plus the brand or an online dealer; C to E to licensed contractors.',
        '2. Suppliers enter the rate for every line (ex-GST), the make and model for section B, and remarks. A line left blank is not priced; do not compare totals until the line counts match.',
        '3. Quantities marked "measure on site" (duct, gas pipe, drain pipe, water pipe) are estimates from the model; the contractor measures and re-prices by the metre.',
        '4. Every quote must also state: steel grade and thickness, price ex-GST with the GST rate, delivery, installation, lead time, payment schedule, warranty and who services it in Goa, validity.',
        '5. Concept drawings for briefing suppliers, not construction drawings.']):
        n.cell(i + 1, 1, t).alignment = Alignment(wrap_text=True, vertical='top')
    n['A1'].font = Font(bold=True, size=13); n.column_dimensions['A'].width = 130
    wb.save(path)

def page():
    def esc(x): return html.escape(str(x))
    body = []
    for r in rows():
        if r[0] == 'sec':
            body.append(f'<tr class="sec"><th>{esc(r[1])}</th><th colspan="7">{esc(r[2])}<small>{esc(r[3])}</small></th></tr>')
        else:
            q = r[6]; q = f'{q:g}' if isinstance(q, float) else q
            body.append(f'<tr><td class="no">{r[1]}</td><td><b>{esc(r[2])}</b></td><td class="sz">{esc(r[3])}</td><td>{esc(r[4])}</td><td>{r[5]}</td><td class="n">{q}</td><td class="blank"></td><td class="blank"></td><td class="ref">{esc(r[7])}</td></tr>')
    n = counts()
    return f'''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover"><title>Kitchen BOQ · Betalbatim Pub</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:opsz,wght@12..96,700&family=IBM+Plex+Sans:wght@400;500;600;700&family=IBM+Plex+Mono:wght@500&display=swap">
<style>
/* Layout: drafting-paper page; the table scrolls sideways on a phone and prints on A4 landscape with a repeating header */
:root{{--paper:#fbf7ef;--ink:#2a2420;--muted:#6a5d4e;--rule:#d9cbb3;--accent:#b23a1e;--chip:#fffdf8;--display:'Bricolage Grotesque','Segoe UI',sans-serif;--body:'IBM Plex Sans','Helvetica Neue',Arial,sans-serif;--mono:'IBM Plex Mono',ui-monospace,Menlo,monospace;color-scheme:light}}
@page{{size:A4 landscape;margin:10mm}}
body{{background:var(--paper);color:var(--ink);font-family:var(--body);margin:0;line-height:1.45}}
.wrap{{max-width:1180px;margin:0 auto;padding:22px 16px 50px;display:flex;flex-direction:column;gap:16px}}
.eyebrow{{font:500 11px/1.4 var(--mono);letter-spacing:.08em;text-transform:uppercase;color:var(--accent);margin:0}}
h1{{font:700 clamp(24px,4vw,34px)/1.05 var(--display);margin:4px 0 6px}}
.lede{{margin:0;max-width:80ch;color:var(--muted)}}
.bar{{display:flex;flex-wrap:wrap;gap:8px}} .bar a{{font:600 13px var(--body);color:var(--ink);background:var(--chip);border:1px solid var(--rule);border-radius:8px;padding:9px 12px;text-decoration:none}} .bar a.pri{{background:var(--ink);color:var(--paper);border-color:var(--ink)}}
.tw{{overflow-x:auto;border:1px solid var(--rule);border-radius:8px;background:var(--chip)}}
table{{border-collapse:collapse;width:100%;min-width:1000px;font-size:13px}}
th,td{{padding:6px 9px;border-top:1px solid var(--rule);vertical-align:top;text-align:left}}
thead th{{background:var(--ink);color:var(--paper);font:600 11.5px var(--mono);border-top:0;white-space:nowrap}}
tr.sec th{{background:#efe6d3;font:700 13px var(--body)}} tr.sec small{{display:block;font:400 12px var(--body);color:var(--muted);margin-top:2px}}
td.no{{font:600 12px var(--mono);white-space:nowrap}} td.sz{{font:500 12px var(--mono);white-space:nowrap}} td.n{{text-align:right;font-family:var(--mono)}} td.ref{{color:var(--muted);font-size:12px}}
td.blank{{background:#fffdf0;min-width:80px}}
ul{{margin:0;padding-left:20px;font-size:14px}} li{{margin:3px 0}}
@media print{{.bar,.hide-print{{display:none}} body{{background:#fff}} .wrap{{padding:0}} .tw{{border:0;overflow:visible}} table{{min-width:0;font-size:8.5pt}} th,td{{padding:3px 6px}} td.blank{{min-width:0;width:11%}} tr.sec small{{font-size:8pt}} td:nth-child(2){{width:22%}} td:nth-child(4){{width:30%}} thead{{display:table-header-group}} tr{{break-inside:avoid}}}}
</style></head><body><div class="wrap">
<header><p class="eyebrow">Simplex · Betalbatim pub · kitchen</p><h1>Bill of quantities (BOQ)</h1>
<p class="lede">{n} priced lines in 8 trades, numbered A.01 to H.06, taken from the model and the two A3 sheets. Each line gives the item or appliance type, its size (W × D × H in mm), and its capacity or rating. Suppliers fill the rate, make and model; the Excel file multiplies and totals. Quantities marked “measure on site” are estimates for pricing. Concept drawings for briefing suppliers, not construction drawings.</p></header>
<div class="bar"><a class="pri" href="plans/kitchen-boq.xlsx">Excel BOQ (rates and totals calculate)</a><a href="plans/kitchen-boq.pdf" target="_blank" rel="noopener">PDF, A4</a><a href="plans/kitchen-boq.csv">CSV</a><a href="boq-goa.html">Equipment and exhaust BOQ (feet, SWG, CFM)</a><a href="kitchen.html">Kitchen sheets</a><a href="quotes.html">Quote comparison</a></div>
<div class="tw"><table><thead><tr><th>Item</th><th>Item / type</th><th>Size W × D × H mm</th><th>Capacity, rating, specification</th><th>Unit</th><th>Qty</th><th>Rate ₹</th><th>Amount ₹</th><th>Drawing</th></tr></thead><tbody>{''.join(body)}</tbody></table></div>
<section><b>Every quote must state</b><ul>
<li>Steel grade and thickness (304 stainless only); make and model of every branded item.</li>
<li>Price ex-GST, itemised, with the GST rate shown separately; delivery to Betalbatim and unloading; installation and connections, included or extra.</li>
<li>Lead time, payment schedule, warranty and who services it in Goa, and how long the quote is valid.</li></ul></section>
</div></body></html>
'''

if __name__ == '__main__':
    write_csv('plans/kitchen-boq.csv'); write_xlsx('plans/kitchen-boq.xlsx')
    open('boq.html', 'w').write(page()); print(counts(), 'lines; boq.html, plans/kitchen-boq.xlsx, .csv written')
